from fastapi import FastAPI, Form, Request
from fastapi.responses import HTMLResponse, RedirectResponse, StreamingResponse
from starlette.middleware.sessions import SessionMiddleware
from datetime import datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

from database import (
    simpan_data,
    update_data,
    ambil_data,
    semua_data,
    data_terbaru,
    hapus_data,
)


app = FastAPI(
    title="E-KESWA PUSBONAT",
    version="1.0.0"
)

app.add_middleware(
    SessionMiddleware,
    secret_key="EKESWA-PUSBONAT-SECRET-2026"
)


USERNAME = "eky"
PASSWORD = "ekynbt26"


def login_required(request: Request):
    return request.session.get("login") is True


def angka(value):
    try:
        return f"{int(value):,}".replace(",", ".")
    except Exception:
        return "0"


def p(value):
    try:
        return f"{float(value):.2f}%"
    except Exception:
        return "0.00%"


CSS = """
<style>

:root{
    --toska:#00a99d;
    --toska-dark:#007f78;
    --toska-light:#e9fbf9;
    --cyan:#00d4c4;
    --navy:#12343b;
    --muted:#72858b;
    --bg:#f3f8f8;
    --white:#ffffff;
    --danger:#e55353;
    --orange:#f59e0b;
    --blue:#3182ce;
    --shadow:0 15px 40px rgba(0,80,75,.09);
}

*{
    box-sizing:border-box;
}

html{
    scroll-behavior:smooth;
}

body{
    margin:0;
    font-family:-apple-system,BlinkMacSystemFont,"Segoe UI",Roboto,Arial,sans-serif;
    background:var(--bg);
    color:var(--navy);
}

a{
    text-decoration:none;
}

button,
input,
select,
textarea{
    font-family:inherit;
}

.app{
    min-height:100vh;
}

.sidebar{
    position:fixed;
    left:0;
    top:0;
    bottom:0;
    width:250px;
    padding:24px 17px;
    background:
        linear-gradient(180deg,#008f86,#006d69);
    color:white;
    z-index:50;
    box-shadow:8px 0 30px rgba(0,80,75,.10);
}

.brand{
    text-align:center;
    padding-bottom:25px;
}

.heart-logo{
    position:relative;
    width:72px;
    height:72px;
    margin:auto;
    border-radius:22px;
    display:flex;
    align-items:center;
    justify-content:center;
    background:rgba(255,255,255,.13);
    border:1px solid rgba(255,255,255,.25);
    box-shadow:
        0 0 25px rgba(0,255,230,.22),
        inset 0 0 20px rgba(255,255,255,.07);
    animation:heartGlow 1.8s infinite;
}

.heart-logo::before{
    content:"♥";
    color:#75fff3;
    font-size:42px;
    text-shadow:
        0 0 8px #00fff0,
        0 0 22px #00d4c4;
}

@keyframes heartGlow{
    0%,100%{
        transform:scale(1);
        box-shadow:0 0 15px rgba(0,255,230,.18);
    }
    50%{
        transform:scale(1.05);
        box-shadow:0 0 32px rgba(0,255,230,.45);
    }
}

.brand h2{
    margin:13px 0 3px;
    font-size:20px;
    letter-spacing:.5px;
}

.brand small{
    opacity:.72;
}

.nav{
    margin-top:18px;
}

.nav a{
    display:flex;
    align-items:center;
    gap:12px;
    color:rgba(255,255,255,.78);
    padding:14px 15px;
    margin:7px 0;
    border-radius:13px;
    font-weight:650;
    transition:.2s;
}

.nav a:hover,
.nav a.active{
    color:#fff;
    background:rgba(255,255,255,.15);
    box-shadow:inset 0 0 0 1px rgba(255,255,255,.08);
}

.nav-icon{
    width:28px;
    text-align:center;
    font-size:18px;
}

.sidebar-bottom{
    position:absolute;
    left:17px;
    right:17px;
    bottom:20px;
}

.logout{
    background:rgba(255,255,255,.10);
}

.main{
    margin-left:250px;
    min-height:100vh;
}

.topbar{
    position:sticky;
    top:0;
    z-index:30;
    padding:17px 30px;
    background:rgba(255,255,255,.86);
    backdrop-filter:blur(18px);
    border-bottom:1px solid rgba(0,100,95,.07);
    display:flex;
    align-items:center;
    justify-content:space-between;
    gap:20px;
}

.topbar-title{
    font-weight:800;
}

.topbar-sub{
    color:var(--muted);
    font-size:13px;
    margin-top:2px;
}

.content{
    padding:30px;
    max-width:1600px;
    margin:auto;
}

.page-title{
    margin-bottom:25px;
}

.page-title h1{
    margin:0;
    font-size:30px;
}

.page-title p{
    color:var(--muted);
    margin:7px 0 0;
}

.cards{
    display:grid;
    grid-template-columns:repeat(4,1fr);
    gap:18px;
    margin-bottom:20px;
}

.card{
    background:rgba(255,255,255,.95);
    border-radius:20px;
    padding:21px;
    box-shadow:var(--shadow);
    border:1px solid rgba(0,140,130,.06);
}

.metric-label{
    color:var(--muted);
    font-size:12px;
    font-weight:750;
    letter-spacing:.7px;
    text-transform:uppercase;
}

.metric{
    font-size:32px;
    font-weight:850;
    color:var(--toska-dark);
    margin-top:7px;
}

.metric-small{
    font-size:13px;
    color:var(--muted);
    margin-top:5px;
}

.grid-2{
    display:grid;
    grid-template-columns:1fr 1fr;
    gap:20px;
}

.grid-3{
    display:grid;
    grid-template-columns:repeat(3,1fr);
    gap:18px;
}

.chart-box{
    height:310px;
}

.progress-wrap{
    margin-top:14px;
}

.progress{
    height:12px;
    border-radius:20px;
    background:#e5eeee;
    overflow:hidden;
}

.progress > div{
    height:100%;
    border-radius:20px;
    background:linear-gradient(90deg,#00a99d,#00ddd0);
    transition:.4s;
}

.percent-big{
    font-size:40px;
    font-weight:900;
    color:var(--toska);
}

.section-title{
    display:flex;
    align-items:center;
    gap:10px;
    margin:0 0 17px;
    font-size:19px;
}

.section-title::before{
    content:"";
    width:5px;
    height:24px;
    border-radius:10px;
    background:var(--toska);
}

.filter{
    display:flex;
    gap:10px;
    flex-wrap:wrap;
    align-items:end;
    margin-bottom:20px;
}

.field{
    flex:1;
    min-width:150px;
}

.field label{
    display:block;
    font-size:12px;
    font-weight:750;
    margin-bottom:7px;
}

input,
select,
textarea{
    width:100%;
    border:1px solid #d9e5e6;
    background:white;
    border-radius:11px;
    padding:12px 13px;
    font-size:14px;
    color:var(--navy);
    outline:none;
}

textarea{
    min-height:100px;
    resize:vertical;
}

input:focus,
select:focus,
textarea:focus{
    border-color:var(--toska);
    box-shadow:0 0 0 3px rgba(0,169,157,.10);
}

.btn{
    border:0;
    border-radius:11px;
    padding:12px 17px;
    cursor:pointer;
    font-weight:750;
    transition:.2s;
    display:inline-flex;
    align-items:center;
    justify-content:center;
    gap:7px;
}

.btn-primary{
    color:white;
    background:linear-gradient(135deg,#00a99d,#007f78);
}

.btn-primary:hover{
    transform:translateY(-1px);
}

.btn-light{
    color:var(--navy);
    background:#edf5f5;
}

.btn-danger{
    color:white;
    background:var(--danger);
}

.btn-excel{
    color:white;
    background:#18864b;
}

.form-section{
    margin-bottom:20px;
}

.form-grid{
    display:grid;
    grid-template-columns:repeat(4,1fr);
    gap:15px;
}

.form-grid-3{
    display:grid;
    grid-template-columns:repeat(3,1fr);
    gap:15px;
}

.form-grid-5{
    display:grid;
    grid-template-columns:repeat(5,1fr);
    gap:15px;
}

.full{
    grid-column:1/-1;
}

.form-title{
    font-size:16px;
    font-weight:800;
    margin-bottom:15px;
}

.auto-percent{
    background:var(--toska-light);
    color:var(--toska-dark);
    border-radius:11px;
    padding:12px;
    font-weight:850;
    text-align:center;
    margin-top:23px;
}

.status-box{
    border-radius:15px;
    padding:15px;
    border:1px solid #e3ecec;
}

.status-rutin{
    background:#eafaf4;
}

.status-tidak{
    background:#fff8e8;
}

.status-putus{
    background:#fff0f0;
}

.status-title{
    font-weight:800;
    margin-bottom:8px;
}

.table-wrap{
    overflow:auto;
}

table{
    width:100%;
    border-collapse:collapse;
    min-width:950px;
}

th,
td{
    padding:12px 10px;
    border-bottom:1px solid #e8eeee;
    text-align:left;
    font-size:13px;
}

th{
    color:#557078;
    background:#f6fafa;
    position:sticky;
    top:0;
}

.action{
    display:flex;
    gap:7px;
}

.badge{
    display:inline-block;
    padding:5px 9px;
    border-radius:8px;
    background:var(--toska-light);
    color:var(--toska-dark);
    font-size:11px;
    font-weight:800;
}

.notice{
    padding:13px 15px;
    border-radius:12px;
    margin-bottom:15px;
    background:#eafaf8;
    color:#006d69;
}

.empty{
    padding:35px;
    text-align:center;
    color:var(--muted);
}

.login-page{
    min-height:100vh;
    display:flex;
    align-items:center;
    justify-content:center;
    padding:20px;
    background:
        radial-gradient(circle at 20% 20%,rgba(0,255,230,.16),transparent 28%),
        radial-gradient(circle at 80% 80%,rgba(0,169,157,.20),transparent 28%),
        linear-gradient(135deg,#003f42,#007f78);
}

.login-glass{
    width:410px;
    max-width:100%;
    padding:38px;
    border-radius:28px;
    color:white;
    background:rgba(255,255,255,.11);
    border:1px solid rgba(255,255,255,.20);
    backdrop-filter:blur(24px);
    box-shadow:0 25px 80px rgba(0,0,0,.25);
}

.login-logo{
    width:90px;
    height:90px;
    margin:0 auto 20px;
    border-radius:28px;
    display:flex;
    align-items:center;
    justify-content:center;
    background:rgba(255,255,255,.12);
    box-shadow:0 0 35px rgba(0,255,230,.30);
    animation:heartGlow 1.8s infinite;
}

.login-logo::before{
    content:"♥";
    color:#6ffff4;
    font-size:55px;
    text-shadow:0 0 12px #00fff0,0 0 30px #00cfc0;
}

.login-glass h1{
    text-align:center;
    margin:0;
    font-size:25px;
}

.login-glass p{
    text-align:center;
    opacity:.72;
    margin:8px 0 28px;
}

.login-glass label{
    display:block;
    margin-bottom:7px;
    font-size:13px;
    font-weight:700;
}

.login-glass input{
    margin-bottom:17px;
    background:rgba(255,255,255,.95);
}

.mobile-menu{
    display:none;
}

@media(max-width:1200px){

    .cards{
        grid-template-columns:repeat(2,1fr);
    }

    .form-grid,
    .form-grid-5{
        grid-template-columns:repeat(2,1fr);
    }
}

@media(max-width:800px){

    .sidebar{
        transform:translateX(-100%);
        transition:.25s;
    }

    .sidebar.open{
        transform:translateX(0);
    }

    .main{
        margin-left:0;
    }

    .mobile-menu{
        display:block;
        border:0;
        background:var(--toska-light);
        border-radius:10px;
        padding:9px 12px;
        font-size:20px;
        cursor:pointer;
    }

    .content{
        padding:20px 15px;
    }

    .topbar{
        padding:13px 15px;
    }

    .cards,
    .grid-2,
    .grid-3,
    .form-grid,
    .form-grid-3,
    .form-grid-5{
        grid-template-columns:1fr;
    }

    .page-title h1{
        font-size:24px;
    }

    .metric{
        font-size:29px;
    }
}

</style>
"""


def shell(body, active="dashboard", title="E-KESWA PUSBONAT"):

    return f"""
<!DOCTYPE html>
<html lang="id">

<head>

<meta charset="UTF-8">

<meta
    name="viewport"
    content="width=device-width, initial-scale=1.0"
>

<title>{title}</title>

{CSS}

</head>

<body>

<div class="app">

<aside class="sidebar" id="sidebar">

<div class="brand">

<div class="heart-logo"></div>

<h2>E-KESWA PUSBONAT</h2>

<small>Sistem Informasi Kesehatan Jiwa</small>

</div>

<nav class="nav">

<a
href="/dashboard"
class="{'active' if active == 'dashboard' else ''}"
>
<span class="nav-icon">📊</span>
<span>Dashboard</span>
</a>

<a
href="/kesehatan-jiwa"
class="{'active' if active == 'keswa' else ''}"
>
<span class="nav-icon">🧠</span>
<span>Kesehatan Jiwa</span>
</a>

</nav>

<div class="sidebar-bottom">

<a href="/logout" class="nav logout">
<span class="nav-icon">↪</span>
<span>Keluar</span>
</a>

</div>

</aside>

<main class="main">

<header class="topbar">

<div>

<button
class="mobile-menu"
onclick="toggleSidebar()"
>
☰
</button>

</div>

<div style="flex:1">

<div class="topbar-title">
{title}
</div>

<div class="topbar-sub">
Puskesmas Pusbónat • Kesehatan Jiwa
</div>

</div>

<div>
<span class="badge">
● ONLINE
</span>
</div>

</header>

<section class="content">

{body}

</section>

</main>

</div>

<script>

function toggleSidebar(){{
    document
        .getElementById("sidebar")
        .classList.toggle("open");
}}

</script>

</body>
</html>
"""


@app.get("/", response_class=HTMLResponse)
async def root(request: Request):

    if login_required(request):
        return RedirectResponse(
            "/dashboard",
            status_code=303
        )

    return login_page()


def login_page():

    return f"""
<!DOCTYPE html>
<html lang="id">

<head>

<meta charset="UTF-8">

<meta
name="viewport"
content="width=device-width,initial-scale=1"
>

<title>E-KESWA PUSBONAT</title>

{CSS}

</head>

<body>

<div class="login-page">

<div class="login-glass">

<div class="login-logo"></div>

<h1>E-KESWA PUSBONAT</h1>

<p>
Sistem Informasi Kesehatan Jiwa
</p>

<form method="post" action="/login">

<label>Username</label>

<input
type="text"
name="username"
placeholder="Masukkan username"
autocomplete="username"
required
>

<label>Password</label>

<input
type="password"
name="password"
placeholder="Masukkan password"
autocomplete="current-password"
required
>

<button
type="submit"
class="btn btn-primary"
style="width:100%;padding:14px"
>
⚡ MASUK APLIKASI
</button>

</form>

<div
style="
text-align:center;
margin-top:22px;
font-size:11px;
opacity:.55;
"
>
E-KESWA • PUSBONAT
</div>

</div>

</div>

</body>
</html>
"""


@app.post("/login")
async def login(
    request: Request,
    username: str = Form(...),
    password: str = Form(...)
):

    if (
        username == USERNAME
        and
        password == PASSWORD
    ):

        request.session["login"] = True

        return RedirectResponse(
            "/dashboard",
            status_code=303
        )

    return HTMLResponse(
        login_page().replace(
            "<p>",
            '<div style="background:rgba(255,100,100,.15);padding:10px;border-radius:10px;text-align:center;margin-bottom:15px">Username atau password salah.</div><p>',
            1
        )
    )


@app.get("/logout")
async def logout(request: Request):

    request.session.clear()

    return RedirectResponse(
        "/",
        status_code=303
    )


def get_filters(request: Request):

    start = request.query_params.get(
        "mulai",
        ""
    )

    end = request.query_params.get(
        "sampai",
        ""
    )

    return start, end


@app.get("/dashboard", response_class=HTMLResponse)
async def dashboard(request: Request):

    if not login_required(request):
        return RedirectResponse("/", status_code=303)

    mulai, sampai = get_filters(request)

    data = data_terbaru(
        mulai if mulai else None,
        sampai if sampai else None
    )

    semua = semua_data(
        mulai if mulai else None,
        sampai if sampai else None
    )

    if data:

        total = data["total_odgj_berat"]

        spm_sasaran = data["spm_sasaran"]
        spm_capaian = data["spm_capaian"]
        spm_persen = data["spm_persentase"]

        f20 = data["f20"]
        f23 = data["f23"]

        desa = [
            data["manusak"],
            data["oelatimo"],
            data["pukdale"],
            data["nunkurus"],
            data["naibonat"],
        ]

        obat_sasaran = data["obat_sasaran"]
        rutin = data["obat_rutin"]
        tidak_rutin = data["obat_tidak_rutin"]
        putus = data["obat_putus"]

        baru = data["odgjb_baru"]

        pel_sasaran = data["pelacakan_sasaran"]
        pel_capaian = data["pelacakan_capaian"]
        pel_persen = data["pelacakan_persentase"]

        p3lp_sasaran = data["p3lp_sasaran"]
        p3lp_capaian = data["p3lp_capaian"]
        p3lp_persen = data["p3lp_persentase"]

        skr_sasaran = data["skrining_sasaran"]
        skr_capaian = data["skrining_capaian"]
        skr_persen = data["skrining_persentase"]

        kendala = data["kendala"]
        solusi = data["solusi"]
        rtl = data["rencana_tindak_lanjut"]

    else:

        total = 0

        spm_sasaran = 0
        spm_capaian = 0
        spm_persen = 0

        f20 = 0
        f23 = 0

        desa = [0,0,0,0,0]

        obat_sasaran = 0
        rutin = 0
        tidak_rutin = 0
        putus = 0

        baru = 0

        pel_sasaran = 0
        pel_capaian = 0
        pel_persen = 0

        p3lp_sasaran = 0
        p3lp_capaian = 0
        p3lp_persen = 0

        skr_sasaran = 0
        skr_capaian = 0
        skr_persen = 0

        kendala = "-"
        solusi = "-"
        rtl = "-"

    max_spm = min(float(spm_persen), 100)
    max_pel = min(float(pel_persen), 100)
    max_p3lp = min(float(p3lp_persen), 100)
    max_skr = min(float(skr_persen), 100)

    periode = (
        f"{mulai or 'Semua'} → {sampai or 'Semua'}"
    )

    body = f"""

<div class="page-title">

<h1>
Dashboard Kesehatan Jiwa
</h1>

<p>
Ringkasan capaian terbaru program Kesehatan Jiwa
</p>

</div>

<form class="filter" method="get" action="/dashboard">

<div class="field">
<label>Tanggal Mulai</label>
<input
type="date"
name="mulai"
value="{mulai}"
>
</div>

<div class="field">
<label>Tanggal Sampai</label>
<input
type="date"
name="sampai"
value="{sampai}"
>
</div>

<button class="btn btn-primary">
🔎 Filter
</button>

<a href="/dashboard" class="btn btn-light">
↻ Reset
</a>

</form>

<div class="notice">
<b>Periode:</b> {periode}
&nbsp; • &nbsp;
<b>Data:</b> {len(semua)} entri
</div>

<div class="cards">

<div class="card">
<div class="metric-label">
Total ODGJ Berat
</div>
<div class="metric">
{angka(total)}
</div>
<div class="metric-small">
Data terbaru
</div>
</div>

<div class="card">
<div class="metric-label">
SPM ODGJ Berat
</div>
<div class="metric">
{angka(spm_capaian)}
</div>
<div class="metric-small">
Sasaran {angka(spm_sasaran)}
</div>
</div>

<div class="card">
<div class="metric-label">
Diagnosis F20
</div>
<div class="metric">
{angka(f20)}
</div>
<div class="metric-small">
Skizofrenia
</div>
</div>

<div class="card">
<div class="metric-label">
Diagnosis F23
</div>
<div class="metric">
{angka(f23)}
</div>
<div class="metric-small">
Gangguan psikotik akut
</div>
</div>

</div>

<div class="grid-2">

<div class="card">

<h3 class="section-title">
SPM Kesehatan Jiwa
</h3>

<div style="display:flex;align-items:end;gap:18px">

<div class="percent-big">
{p(spm_persen)}
</div>

<div style="padding-bottom:7px;color:var(--muted)">
{angka(spm_capaian)} / {angka(spm_sasaran)}
</div>

</div>

<div class="progress-wrap">
<div class="progress">
<div style="width:{max_spm}%"></div>
</div>
</div>

</div>

<div class="card">

<h3 class="section-title">
Pelacakan ODGJB
</h3>

<div style="display:flex;align-items:end;gap:18px">

<div class="percent-big">
{p(pel_persen)}
</div>

<div style="padding-bottom:7px;color:var(--muted)">
{angka(pel_capaian)} / {angka(pel_sasaran)}
</div>

</div>

<div class="progress-wrap">
<div class="progress">
<div style="width:{max_pel}%"></div>
</div>
</div>

</div>

</div>

<br>

<div class="grid-2">

<div class="card">

<h3 class="section-title">
Sebaran ODGJB per Desa/Kelurahan
</h3>

<div class="chart-box">
<canvas id="desaChart"></canvas>
</div>

</div>

<div class="card">

<h3 class="section-title">
Status Pemantauan Obat
</h3>

<div class="chart-box">
<canvas id="obatChart"></canvas>
</div>

<div
style="
display:flex;
justify-content:space-between;
margin-top:10px;
font-size:13px;
color:var(--muted)
"
>
<span>
Sasaran: <b>{angka(obat_sasaran)}</b>
</span>

<span>
Terpantau: <b>{angka(rutin + tidak_rutin + putus)}</b>
</span>
</div>

</div>

</div>

<br>

<div class="cards">

<div class="card">

<div class="metric-label">
ODGJB Baru
</div>

<div class="metric">
{angka(baru)}
</div>

</div>

<div class="card">

<div class="metric-label">
P3LP
</div>

<div class="metric">
{angka(p3lp_capaian)}
</div>

<div class="metric-small">
Sasaran {angka(p3lp_sasaran)}
• {p(p3lp_persen)}
</div>

</div>

<div class="card">

<div class="metric-label">
Skrining ODGJB
</div>

<div class="metric">
{angka(skr_capaian)}
</div>

<div class="metric-small">
Sasaran {angka(skr_sasaran)}
• {p(skr_persen)}
</div>

</div>

<div class="card">

<div class="metric-label">
Entri Data
</div>

<div class="metric">
{angka(len(semua))}
</div>

<div class="metric-small">
Riwayat tersimpan
</div>

</div>

</div>

<div class="grid-2">

<div class="card">

<h3 class="section-title">
Skrining ODGJB
</h3>

<div class="percent-big">
{p(skr_persen)}
</div>

<p>
Capaian {angka(skr_capaian)}
dari sasaran {angka(skr_sasaran)}
</p>

<div class="progress">
<div style="width:{max_skr}%"></div>
</div>

</div>

<div class="card">

<h3 class="section-title">
P3LP
</h3>

<div class="percent-big">
{p(p3lp_persen)}
</div>

<p>
Capaian {angka(p3lp_capaian)}
dari sasaran {angka(p3lp_sasaran)}
</p>

<div class="progress">
<div style="width:{max_p3lp}%"></div>
</div>

</div>

</div>

<br>

<div class="card">

<h3 class="section-title">
Kendala, Solusi & Rencana Tindak Lanjut
</h3>

<div class="grid-3">

<div>
<b>Kendala</b>
<p style="color:var(--muted);white-space:pre-line">
{kendala}
</p>
</div>

<div>
<b>Solusi</b>
<p style="color:var(--muted);white-space:pre-line">
{solusi}
</p>
</div>

<div>
<b>Rencana Tindak Lanjut</b>
<p style="color:var(--muted);white-space:pre-line">
{rtl}
</p>
</div>

</div>

</div>

<br>

<div class="card">

<h3 class="section-title">
Riwayat Entri
</h3>

<div class="table-wrap">

<table>

<tr>
<th>Tanggal</th>
<th>Total ODGJB</th>
<th>SPM</th>
<th>F20</th>
<th>F23</th>
<th>Pelacakan</th>
<th>Skrining</th>
<th>Aksi</th>
</tr>

"""

    if semua:

        for row in semua[:30]:

            body += f"""

<tr>

<td>
{row["tanggal"]}
</td>

<td>
{angka(row["total_odgj_berat"])}
</td>

<td>
{angka(row["spm_capaian"])}
/
{angka(row["spm_sasaran"])}
<br>
<span class="badge">
{p(row["spm_persentase"])}
</span>
</td>

<td>
{angka(row["f20"])}
</td>

<td>
{angka(row["f23"])}
</td>

<td>
{angka(row["pelacakan_capaian"])}
/
{angka(row["pelacakan_sasaran"])}
</td>

<td>
{angka(row["skrining_capaian"])}
/
{angka(row["skrining_sasaran"])}
</td>

<td>

<div class="action">

<a
class="btn btn-light"
href="/kesehatan-jiwa/edit/{row["id"]}"
>
✏️ Edit
</a>

<form
method="post"
action="/hapus/{row["id"]}"
onsubmit="return confirm('Hapus data ini?')"
>
<button
class="btn btn-danger"
type="submit"
>
🗑
</button>
</form>

</div>

</td>

</tr>

"""

    else:

        body += """
<tr>
<td colspan="8">
<div class="empty">
Belum ada data pada periode ini.
</div>
</td>
</tr>
"""

    body += """

</table>

</div>

</div>

<script src="https://cdn.jsdelivr.net/npm/chart.js"></script>

<script>

new Chart(
document.getElementById("desaChart"),
{
type:"bar",

data:{
labels:[
"Manusak",
"Oelatimo",
"Pukdale",
"Nunkurus",
"Naibonat"
],

datasets:[
{
label:"Jumlah ODGJB",
data:[
""" + ",".join(str(x) for x in desa) + """
],
borderRadius:8
}
]
},

options:{
responsive:true,
maintainAspectRatio:false,

plugins:{
legend:{
display:false
}
},

scales:{
y:{
beginAtZero:true
}
}
}
}
);


new Chart(
document.getElementById("obatChart"),
{
type:"doughnut",

data:{
labels:[
"Rutin",
"Tidak Rutin",
"Putus"
],

datasets:[
{
data:[
""" + str(rutin) + "," + str(tidak_rutin) + "," + str(putus) + """
],
borderWidth:0
}
]
},

options:{
responsive:true,
maintainAspectRatio:false,
cutout:"65%"
}
}
);


/* AUTO REFRESH DASHBOARD
   setiap 3 detik agar data terbaru
   muncul tanpa harus refresh manual.
*/

setTimeout(function(){

    const url =
        window.location.href;

    fetch(url,{
        headers:{
            "X-Requested-With":"XMLHttpRequest"
        }
    })
    .then(function(response){

        if(response.ok){

            window.location.reload();

        }

    })
    .catch(function(){});

},3000);

</script>

"""

    return HTMLResponse(
        shell(
            body,
            active="dashboard",
            title="Dashboard"
        )
    )


def form_page(data=None, edit_id=None):

    is_edit = data is not None

    title = (
        "Edit Data Kesehatan Jiwa"
        if is_edit
        else
        "Input Kesehatan Jiwa"
    )

    action = (
        f"/kesehatan-jiwa/edit/{edit_id}"
        if is_edit
        else
        "/kesehatan-jiwa"
    )

    def val(key):
        if data:
            return data[key] if data[key] is not None else ""
        return ""

    body = f"""

<div class="page-title">

<h1>
{title}
</h1>

<p>
Masukkan data program Kesehatan Jiwa.
Semua persentase dihitung otomatis.
</p>

</div>

<form
method="post"
action="{action}"
>

<div class="card form-section">

<div class="form-title">
📅 Periode Data
</div>

<div class="form-grid">

<div class="field">

<label>
Tanggal
</label>

<input
type="date"
name="tanggal"
value="{val("tanggal")}"
required
>

</div>

</div>

</div>


<div class="card form-section">

<div class="form-title">
🧠 Data Utama ODGJ Berat
</div>

<div class="form-grid-3">

<div class="field">

<label>
Total ODGJ Berat
</label>

<input
type="number"
name="total_odgj_berat"
min="0"
value="{val("total_odgj_berat")}"
required
>

</div>

<div class="field">

<label>
Diagnosis F20
</label>

<input
type="number"
name="f20"
min="0"
value="{val("f20")}"
>

</div>

<div class="field">

<label>
Diagnosis F23
</label>

<input
type="number"
name="f23"
min="0"
value="{val("f23")}"
>

</div>

</div>

</div>


<div class="card form-section">

<div class="form-title">
📊 SPM ODGJ Berat
</div>

<div class="form-grid-3">

<div class="field">

<label>
Sasaran SPM
</label>

<input
id="spm_sasaran"
type="number"
name="spm_sasaran"
min="0"
value="{val("spm_sasaran")}"
oninput="hitungPersen('spm')"
>

</div>

<div class="field">

<label>
Capaian SPM
</label>

<input
id="spm_capaian"
type="number"
name="spm_capaian"
min="0"
value="{val("spm_capaian")}"
oninput="hitungPersen('spm')"
>

</div>

<div
class="auto-percent"
id="spm_hasil"
>
0.00%
</div>

</div>

</div>


<div class="card form-section">

<div class="form-title">
📍 Sebaran ODGJB Desa/Kelurahan
</div>

<div class="form-grid-5">

<div class="field">
<label>Desa Manusak</label>
<input
type="number"
name="manusak"
min="0"
value="{val("manusak")}"
>
</div>

<div class="field">
<label>Desa Oelatimo</label>
<input
type="number"
name="oelatimo"
min="0"
value="{val("oelatimo")}"
>
</div>

<div class="field">
<label>Desa Pukdale</label>
<input
type="number"
name="pukdale"
min="0"
value="{val("pukdale")}"
>
</div>

<div class="field">
<label>Desa Nunkurus</label>
<input
type="number"
name="nunkurus"
min="0"
value="{val("nunkurus")}"
>
</div>

<div class="field">
<label>Kelurahan Naibonat</label>
<input
type="number"
name="naibonat"
min="0"
value="{val("naibonat")}"
>
</div>

</div>

</div>


<div class="card form-section">

<div class="form-title">
💊 Pemantauan Obat ODGJB
</div>

<div class="form-grid">

<div class="field">

<label>
Sasaran Pemantauan
</label>

<input
type="number"
name="obat_sasaran"
min="0"
value="{val("obat_sasaran")}"
>

</div>

<div class="status-box status-rutin">

<div class="status-title">
🟢 Rutin
</div>

<input
type="number"
name="obat_rutin"
min="0"
value="{val("obat_rutin")}"
>

</div>

<div class="status-box status-tidak">

<div class="status-title">
🟡 Tidak Rutin
</div>

<input
type="number"
name="obat_tidak_rutin"
min="0"
value="{val("obat_tidak_rutin")}"
>

</div>

<div class="status-box status-putus">

<div class="status-title">
🔴 Putus
</div>

<input
type="number"
name="obat_putus"
min="0"
value="{val("obat_putus")}"
>

</div>

</div>

</div>


<div class="card form-section">

<div class="form-title">
🆕 ODGJB Baru
</div>

<div class="form-grid">

<div class="field">

<label>
Jumlah ODGJB Baru
</label>

<input
type="number"
name="odgjb_baru"
min="0"
value="{val("odgjb_baru")}"
>

</div>

</div>

</div>


<div class="card form-section">

<div class="form-title">
🔎 Pelacakan ODGJB
</div>

<div class="form-grid-3">

<div class="field">

<label>
Sasaran Pelacakan
</label>

<input
id="pel_sasaran"
type="number"
name="pelacakan_sasaran"
min="0"
value="{val("pelacakan_sasaran")}"
oninput="hitungPersen('pel')"
>

</div>

<div class="field">

<label>
Capaian Pelacakan
</label>

<input
id="pel_capaian"
type="number"
name="pelacakan_capaian"
min="0"
value="{val("pelacakan_capaian")}"
oninput="hitungPersen('pel')"
>

</div>

<div
class="auto-percent"
id="pel_hasil"
>
0.00%
</div>

</div>

</div>


<div class="card form-section">

<div class="form-title">
🩺 P3LP
</div>

<div class="form-grid-3">

<div class="field">

<label>
Sasaran P3LP
</label>

<input
id="p3lp_sasaran"
type="number"
name="p3lp_sasaran"
min="0"
value="{val("p3lp_sasaran")}"
oninput="hitungPersen('p3lp')"
>

</div>

<div class="field">

<label>
Capaian P3LP
</label>

<input
id="p3lp_capaian"
type="number"
name="p3lp_capaian"
min="0"
value="{val("p3lp_capaian")}"
oninput="hitungPersen('p3lp')"
>

</div>

<div
class="auto-percent"
id="p3lp_hasil"
>
0.00%
</div>

</div>

</div>


<div class="card form-section">

<div class="form-title">
🔬 Skrining ODGJB
</div>

<div class="form-grid-3">

<div class="field">

<label>
Sasaran Skrining
</label>

<input
id="skr_sasaran"
type="number"
name="skrining_sasaran"
min="0"
value="{val("skrining_sasaran")}"
oninput="hitungPersen('skr')"
>

</div>

<div class="field">

<label>
Capaian Skrining
</label>

<input
id="skr_capaian"
type="number"
name="skrining_capaian"
min="0"
value="{val("skrining_capaian")}"
oninput="hitungPersen('skr')"
>

</div>

<div
class="auto-percent"
id="skr_hasil"
>
0.00%
</div>

</div>

</div>


<div class="card form-section">

<div class="form-title">
📝 Analisis Program
</div>

<div class="form-grid-3">

<div class="field">

<label>
Kendala
</label>

<textarea
name="kendala"
placeholder="Tuliskan kendala yang ditemukan..."
>{val("kendala")}</textarea>

</div>

<div class="field">

<label>
Solusi
</label>

<textarea
name="solusi"
placeholder="Tuliskan solusi yang dilakukan..."
>{val("solusi")}</textarea>

</div>

<div class="field">

<label>
Rencana Tindak Lanjut
</label>

<textarea
name="rencana_tindak_lanjut"
placeholder="Tuliskan rencana tindak lanjut..."
>{val("rencana_tindak_lanjut")}</textarea>

</div>

</div>

</div>


<div
style="
display:flex;
gap:10px;
justify-content:flex-end;
flex-wrap:wrap;
margin-bottom:30px;
"
>

<a
href="/dashboard"
class="btn btn-light"
>
Batal
</a>

<button
type="submit"
class="btn btn-primary"
>
💾 Simpan Data
</button>

</div>

</form>


<script>

function hitungPersen(prefix){{

    let sasaran;
    let capaian;

    if(prefix==="spm"){{
        sasaran =
            Number(
                document.getElementById(
                    "spm_sasaran"
                ).value
            );

        capaian =
            Number(
                document.getElementById(
                    "spm_capaian"
                ).value
            );
    }}

    if(prefix==="pel"){{
        sasaran =
            Number(
                document.getElementById(
                    "pel_sasaran"
                ).value
            );

        capaian =
            Number(
                document.getElementById(
                    "pel_capaian"
                ).value
            );
    }}

    if(prefix==="p3lp"){{
        sasaran =
            Number(
                document.getElementById(
                    "p3lp_sasaran"
                ).value
            );

        capaian =
            Number(
                document.getElementById(
                    "p3lp_capaian"
                ).value
            );
    }}

    if(prefix==="skr"){{
        sasaran =
            Number(
                document.getElementById(
                    "skr_sasaran"
                ).value
            );

        capaian =
            Number(
                document.getElementById(
                    "skr_capaian"
                ).value
            );
    }}

    let hasil = 0;

    if(sasaran > 0){{
        hasil =
            (capaian / sasaran) * 100;
    }}

    document.getElementById(
        prefix + "_hasil"
    ).innerText =
        hasil.toFixed(2) + "%";
}}


hitungPersen("spm");
hitungPersen("pel");
hitungPersen("p3lp");
hitungPersen("skr");

</script>

"""

    return HTMLResponse(
        shell(
            body,
            active="keswa",
            title="Kesehatan Jiwa"
        )
    )


def ambil_form(
    tanggal,
    total_odgj_berat,
    spm_sasaran,
    spm_capaian,
    f23,
    f20,
    manusak,
    oelatimo,
    pukdale,
    nunkurus,
    naibonat,
    obat_sasaran,
    obat_rutin,
    obat_tidak_rutin,
    obat_putus,
    odgjb_baru,
    pelacakan_sasaran,
    pelacakan_capaian,
    p3lp_sasaran,
    p3lp_capaian,
    skrining_sasaran,
    skrining_capaian,
    kendala,
    solusi,
    rencana_tindak_lanjut,
):

    return {

        "tanggal": tanggal,

        "total_odgj_berat":
            int(total_odgj_berat or 0),

        "spm_sasaran":
            int(spm_sasaran or 0),

        "spm_capaian":
            int(spm_capaian or 0),

        "f23":
            int(f23 or 0),

        "f20":
            int(f20 or 0),

        "manusak":
            int(manusak or 0),

        "oelatimo":
            int(oelatimo or 0),

        "pukdale":
            int(pukdale or 0),

        "nunkurus":
            int(nunkurus or 0),

        "naibonat":
            int(naibonat or 0),

        "obat_sasaran":
            int(obat_sasaran or 0),

        "obat_rutin":
            int(obat_rutin or 0),

        "obat_tidak_rutin":
            int(obat_tidak_rutin or 0),

        "obat_putus":
            int(obat_putus or 0),

        "odgjb_baru":
            int(odgjb_baru or 0),

        "pelacakan_sasaran":
            int(pelacakan_sasaran or 0),

        "pelacakan_capaian":
            int(pelacakan_capaian or 0),

        "p3lp_sasaran":
            int(p3lp_sasaran or 0),

        "p3lp_capaian":
            int(p3lp_capaian or 0),

        "skrining_sasaran":
            int(skrining_sasaran or 0),

        "skrining_capaian":
            int(skrining_capaian or 0),

        "kendala":
            kendala or "",

        "solusi":
            solusi or "",

        "rencana_tindak_lanjut":
            rencana_tindak_lanjut or "",
    }


@app.get(
    "/kesehatan-jiwa",
    response_class=HTMLResponse
)
async def kesehatan_jiwa(request: Request):

    if not login_required(request):
        return RedirectResponse("/", status_code=303)

    return form_page()


@app.post("/kesehatan-jiwa")
async def simpan_keswa(
    request: Request,

    tanggal: str = Form(...),

    total_odgj_berat: int = Form(0),

    spm_sasaran: int = Form(0),
    spm_capaian: int = Form(0),

    f23: int = Form(0),
    f20: int = Form(0),

    manusak: int = Form(0),
    oelatimo: int = Form(0),
    pukdale: int = Form(0),
    nunkurus: int = Form(0),
    naibonat: int = Form(0),

    obat_sasaran: int = Form(0),
    obat_rutin: int = Form(0),
    obat_tidak_rutin: int = Form(0),
    obat_putus: int = Form(0),

    odgjb_baru: int = Form(0),

    pelacakan_sasaran: int = Form(0),
    pelacakan_capaian: int = Form(0),

    p3lp_sasaran: int = Form(0),
    p3lp_capaian: int = Form(0),

    skrining_sasaran: int = Form(0),
    skrining_capaian: int = Form(0),

    kendala: str = Form(""),
    solusi: str = Form(""),
    rencana_tindak_lanjut: str = Form(""),
):

    if not login_required(request):
        return RedirectResponse("/", status_code=303)

    data = ambil_form(
        tanggal,
        total_odgj_berat,

        spm_sasaran,
        spm_capaian,

        f23,
        f20,

        manusak,
        oelatimo,
        pukdale,
        nunkurus,
        naibonat,

        obat_sasaran,
        obat_rutin,
        obat_tidak_rutin,
        obat_putus,

        odgjb_baru,

        pelacakan_sasaran,
        pelacakan_capaian,

        p3lp_sasaran,
        p3lp_capaian,

        skrining_sasaran,
        skrining_capaian,

        kendala,
        solusi,
        rencana_tindak_lanjut,
    )

    simpan_data(data)

    return RedirectResponse(
        "/dashboard",
        status_code=303
    )


@app.get(
    "/kesehatan-jiwa/edit/{record_id}",
    response_class=HTMLResponse
)
async def edit_page(
    request: Request,
    record_id: int
):

    if not login_required(request):
        return RedirectResponse("/", status_code=303)

    data = ambil_data(record_id)

    if not data:
        return RedirectResponse(
            "/kesehatan-jiwa",
            status_code=303
        )

    return form_page(
        data=data,
        edit_id=record_id
    )


@app.post("/kesehatan-jiwa/edit/{record_id}")
async def update_keswa(
    request: Request,
    record_id: int,

    tanggal: str = Form(...),

    total_odgj_berat: int = Form(0),

    spm_sasaran: int = Form(0),
    spm_capaian: int = Form(0),

    f23: int = Form(0),
    f20: int = Form(0),

    manusak: int = Form(0),
    oelatimo: int = Form(0),
    pukdale: int = Form(0),
    nunkurus: int = Form(0),
    naibonat: int = Form(0),

    obat_sasaran: int = Form(0),
    obat_rutin: int = Form(0),
    obat_tidak_rutin: int = Form(0),
    obat_putus: int = Form(0),

    odgjb_baru: int = Form(0),

    pelacakan_sasaran: int = Form(0),
    pelacakan_capaian: int = Form(0),

    p3lp_sasaran: int = Form(0),
    p3lp_capaian: int = Form(0),

    skrining_sasaran: int = Form(0),
    skrining_capaian: int = Form(0),

    kendala: str = Form(""),
    solusi: str = Form(""),
    rencana_tindak_lanjut: str = Form(""),
):

    if not login_required(request):
        return RedirectResponse("/", status_code=303)

    data = ambil_form(
        tanggal,
        total_odgj_berat,

        spm_sasaran,
        spm_capaian,

        f23,
        f20,

        manusak,
        oelatimo,
        pukdale,
        nunkurus,
        naibonat,

        obat_sasaran,
        obat_rutin,
        obat_tidak_rutin,
        obat_putus,

        odgjb_baru,

        pelacakan_sasaran,
        pelacakan_capaian,

        p3lp_sasaran,
        p3lp_capaian,

        skrining_sasaran,
        skrining_capaian,

        kendala,
        solusi,
        rencana_tindak_lanjut,
    )

    update_data(
        record_id,
        data
    )

    return RedirectResponse(
        "/dashboard",
        status_code=303
    )


@app.post("/hapus/{record_id}")
async def delete_data(
    request: Request,
    record_id: int
):

    if not login_required(request):
        return RedirectResponse("/", status_code=303)

    hapus_data(record_id)

    return RedirectResponse(
        "/dashboard",
        status_code=303
    )


@app.get("/export")
async def export_excel(request: Request):

    if not login_required(request):
        return RedirectResponse("/", status_code=303)

    mulai, sampai = get_filters(request)

    rows = semua_data(
        mulai if mulai else None,
        sampai if sampai else None
    )

    wb = Workbook()

    ws = wb.active

    ws.title = "E-KESWA"

    headers = [
        "No",
        "Tanggal",
        "Total ODGJ Berat",

        "SPM Sasaran",
        "SPM Capaian",
        "SPM Persentase",

        "F23",
        "F20",

        "Manusak",
        "Oelatimo",
        "Pukdale",
        "Nunkurus",
        "Naibonat",

        "Obat Sasaran",
        "Obat Rutin",
        "Obat Tidak Rutin",
        "Obat Putus",

        "ODGJB Baru",

        "Pelacakan Sasaran",
        "Pelacakan Capaian",
        "Pelacakan Persentase",

        "P3LP Sasaran",
        "P3LP Capaian",
        "P3LP Persentase",

        "Skrining Sasaran",
        "Skrining Capaian",
        "Skrining Persentase",

        "Kendala",
        "Solusi",
        "Rencana Tindak Lanjut",
    ]

    ws.append(headers)

    for cell in ws[1]:

        cell.font = Font(
            bold=True,
            color="FFFFFF"
        )

        cell.fill = PatternFill(
            "solid",
            fgColor="008F86"
        )

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )

    for no, row in enumerate(
        rows,
        start=1
    ):

        ws.append([

            no,

            row["tanggal"],
            row["total_odgj_berat"],

            row["spm_sasaran"],
            row["spm_capaian"],
            row["spm_persentase"],

            row["f23"],
            row["f20"],

            row["manusak"],
            row["oelatimo"],
            row["pukdale"],
            row["nunkurus"],
            row["naibonat"],

            row["obat_sasaran"],
            row["obat_rutin"],
            row["obat_tidak_rutin"],
            row["obat_putus"],

            row["odgjb_baru"],

            row["pelacakan_sasaran"],
            row["pelacakan_capaian"],
            row["pelacakan_persentase"],

            row["p3lp_sasaran"],
            row["p3lp_capaian"],
            row["p3lp_persentase"],

            row["skrining_sasaran"],
            row["skrining_capaian"],
            row["skrining_persentase"],

            row["kendala"],
            row["solusi"],
            row["rencana_tindak_lanjut"],
        ])

    widths = [
        7,18,18,
        15,15,18,
        10,10,
        12,12,12,12,15,
        15,15,18,15,
        15,
        18,18,22,
        15,15,20,
        18,18,22,
        35,35,40
    ]

    for index, width in enumerate(
        widths,
        start=1
    ):

        ws.column_dimensions[
            ws.cell(1,index).column_letter
        ].width = width

    ws.freeze_panes = "A2"

    output = BytesIO()

    wb.save(output)

    output.seek(0)

    filename = (
        "E-KESWA_PUSBONAT_"
        +
        datetime.now().strftime(
            "%Y%m%d_%H%M%S"
        )
        +
        ".xlsx"
    )

    return StreamingResponse(
        output,
        media_type=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
        headers={
            "Content-Disposition":
            f'attachment; filename="{filename}"'
        }
    )


@app.get("/api/dashboard")
async def api_dashboard(
    request: Request
):

    if not login_required(request):
        return {
            "login": False
        }

    mulai, sampai = get_filters(request)

    data = data_terbaru(
        mulai if mulai else None,
        sampai if sampai else None
    )

    if not data:
        return {
            "login": True,
            "data": None
        }

    return {
        "login": True,
        "data": dict(data)
    }


if __name__ == "__main__":

    import uvicorn

    uvicorn.run(
        "main:app",
        host="127.0.0.1",
        port=8010,
        reload=False
    )
